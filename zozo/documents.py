"""بایگانی اسناد: دریافت فایل، استخراج متن (با OCR فارسی در صورت نیاز)، نمایه‌سازی و جستجو. برگرفته از CivilDesk.

مسیر پردازش هر صفحه‌ی PDF:
1. استخراج لایه‌ی متنی ← درست کردن حروف شکل‌دار و متن برعکس
2. سنجش کیفیت؛ اگر صفحه اسکن‌شده، کم‌متن یا خراب بود ← OCR فارسی/انگلیسی با Tesseract
3. تقسیم به بندهای کوتاه با شماره‌ی صفحه و نمایه‌ی کلیدواژه‌ای (FTS5)

فایل‌های صوتی و ویدیویی (مثلاً ضبط مصاحبه) هم بایگانی می‌شوند؛ جستجو در عنوان، برچسب و یادداشت آن‌هاست.
(جای خالی برای آینده: اگر مدل تبدیل گفتار به متن در دسترس بود، در process_document به متن تبدیل شوند.)
"""
from __future__ import annotations

import hashlib
import logging
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile
import threading
from pathlib import Path
from typing import Any, BinaryIO, Iterator

from . import db, jalali, textnorm
from .config import settings

log = logging.getLogger(__name__)

PDF_EXT = {".pdf"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
TEXT_EXT = {".txt", ".md"}
DOCX_EXT = {".docx"}
AUDIO_EXT = {".mp3", ".m4a", ".wav", ".ogg", ".oga", ".opus", ".aac", ".amr", ".flac"}
VIDEO_EXT = {".mp4", ".mov", ".mkv", ".webm", ".avi", ".3gp"}
OTHER_EXT = {".xlsx", ".xls", ".pptx", ".zip", ".rar", ".csv", ".doc"}

CATEGORIES = {
    "document": "سند / مدرک", "letter": "نامه و مکاتبه", "contract": "قرارداد", "interview": "مصاحبه",
    "report": "گزارش / پژوهش", "press": "بیانیه / خبرنامه", "photo": "عکس", "audio": "صوت", "video": "ویدیو",
    "personal": "مدارک شخصی", "other": "سایر",
}

CHUNK_CHARS = 900
CHUNK_OVERLAP = 150
OCR_DPI = 300
OCR_MAX_PIXELS = 4200
GOOD_SCORE = 0.97


# کارهای سنگین (OCR، ساخت ویدیو، تبدیل گفتار) نوبتی اجرا می‌شوند تا حافظه‌ی سرور کوچک پر نشود
HEAVY = threading.Lock()


def docs_dir() -> Path:
    return settings.dir("docs")


# ───────────────────────── OCR ─────────────────────────


def ocr_available() -> bool:
    return shutil.which("tesseract") is not None and os.getenv("ZOZO_NO_OCR") is None


def _tesseract(path: str, langs: str) -> str:
    res = subprocess.run(
        ["tesseract", path, "stdout", "-l", langs, "--psm", "3"],
        capture_output=True, text=True, timeout=300,
    )
    if res.returncode != 0:
        raise RuntimeError(res.stderr.strip()[:300] or "خطای OCR")
    return res.stdout


def ocr_image(path: str) -> str:
    """OCR تطبیقی: اول فارسی، اگر نتیجه ضعیف بود انگلیسی و ترکیبی."""
    first = os.getenv("ZOZO_OCR_LANGS", "fas")
    best, best_q = "", -1.0
    for langs in dict.fromkeys([first, "eng", "fas+eng"]):
        text = _tesseract(path, langs)
        qd = textnorm.quality(text)
        q = qd["score"] * min(1.0, 0.5 + qd["words"] / 40)
        if q > best_q:
            best, best_q = text, q
        if qd["score"] >= 0.85 and qd["words"] >= 15:
            break
    return best


def _ocr_pdf_page(page) -> str:
    import pymupdf

    longest = max(page.rect.width, page.rect.height) / 72
    dpi = min(OCR_DPI, int(OCR_MAX_PIXELS / max(longest, 1)))
    pix = page.get_pixmap(dpi=max(dpi, 120), colorspace=pymupdf.csGRAY)
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        pix.save(tmp.name)
    try:
        return ocr_image(tmp.name)
    finally:
        os.unlink(tmp.name)


# ───────────────────────── استخراج متن ─────────────────────────


def _best_text(text_layer: str, ocr_fn) -> tuple[str, str]:
    fixed, reversed_ = textnorm.fix_direction(text_layer)
    q = textnorm.quality(fixed)
    method = "fixed" if reversed_ else "text"
    if q["words"] >= 15 and q["score"] >= GOOD_SCORE and q["bad"] <= 0.02:
        return fixed, method
    if ocr_fn is None:
        return fixed, (method if q["words"] >= 15 and q["score"] >= 0.6 else "weak")
    try:
        ocr = ocr_fn()
    except Exception as e:
        log.warning("OCR ناموفق: %s", e)
        return fixed, "weak"
    oq = textnorm.quality(ocr)
    if oq["words"] > 0 and (oq["score"] > q["score"] + 0.02 or (q["words"] < 15 and oq["words"] > q["words"])):
        return ocr, "ocr"
    return fixed, method if q["score"] >= 0.6 else "weak"


def iter_pages(path: Path, kind: str) -> Iterator[tuple[int, int, Any]]:
    """(شماره‌ی صفحه، تعداد کل، تابع استخراج) برای هر صفحه."""
    use_ocr = ocr_available()
    if kind == "pdf":
        import pymupdf

        doc = pymupdf.open(path)
        try:
            n = doc.page_count
            for i in range(n):
                page = doc[i]
                yield i + 1, n, lambda page=page: _best_text(
                    page.get_text("text"), (lambda: _ocr_pdf_page(page)) if use_ocr else None
                )
        finally:
            doc.close()
    elif kind == "image":
        if not use_ocr:
            yield 1, 1, lambda: ("", "weak")
            return
        yield 1, 1, lambda: (ocr_image(str(path)), "ocr")
    elif kind == "docx":
        import docx

        d = docx.Document(str(path))
        paras = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                paras.append(" | ".join(c.text for c in row.cells))
        yield from _virtual_pages("\n".join(paras))
    elif kind == "text":
        raw = path.read_bytes()
        for enc in ("utf-8", "utf-16", "cp1256"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = raw.decode("utf-8", "replace")
        yield from _virtual_pages(text)
    # صوت، ویدیو و سایر: متنی برای استخراج ندارند


def _virtual_pages(text: str, size: int = 3000) -> Iterator[tuple[int, int, Any]]:
    parts, cur = [], ""
    for para in text.split("\n"):
        if len(cur) + len(para) > size and cur:
            parts.append(cur)
            cur = ""
        cur += para + "\n"
    if cur.strip():
        parts.append(cur)
    parts = parts or [""]
    for i, p in enumerate(parts):
        yield i + 1, len(parts), lambda p=p: (p, "text")


# ───────────────────────── تقسیم به بند ─────────────────────────

_SENT_END = re.compile(r"(?<=[.!?؟؛:])\s+")


def _join_lines(text: str) -> list[str]:
    paras, cur = [], []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            if cur:
                paras.append(" ".join(cur))
                cur = []
            continue
        cur.append(line)
        if line.endswith((".", "؟", "?", ":", "!")) and len(" ".join(cur)) > 200:
            paras.append(" ".join(cur))
            cur = []
    if cur:
        paras.append(" ".join(cur))
    return [p for p in paras if p]


def chunk_text(text: str) -> list[str]:
    text = textnorm.clean_display(text)
    chunks, cur = [], ""
    for para in _join_lines(text):
        pieces = [para] if len(para) <= CHUNK_CHARS else [s for s in _SENT_END.split(para) if s]
        for piece in pieces:
            if len(piece) > CHUNK_CHARS and cur:
                chunks.append(cur.strip())
                cur = ""
            while len(piece) > CHUNK_CHARS:
                chunks.append(piece[:CHUNK_CHARS].strip())
                piece = piece[CHUNK_CHARS - CHUNK_OVERLAP:]
            if len(cur) + len(piece) + 1 > CHUNK_CHARS and cur:
                chunks.append(cur.strip())
                tail = cur[-CHUNK_OVERLAP:]
                cur = tail[tail.find(" ") + 1:] if " " in tail else ""
            cur += ("\n" if cur and piece is para else " ") + piece
    if cur.strip():
        chunks.append(cur.strip())
    return [c for c in chunks if len(textnorm.normalize(c)) >= 3]


# ───────────────────────── دریافت فایل ─────────────────────────


def kind_of(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    for kind, exts in (("pdf", PDF_EXT), ("image", IMAGE_EXT), ("docx", DOCX_EXT), ("text", TEXT_EXT),
                       ("audio", AUDIO_EXT), ("video", VIDEO_EXT), ("other", OTHER_EXT)):
        if ext in exts:
            return kind
    raise db.ValidationError(
        f"نوع فایل پشتیبانی نمی‌شود: {ext or filename}. مجاز: PDF، عکس، Word، متن، صوت، ویدیو، Excel"
    )


def default_category(kind: str) -> str:
    return {"image": "photo", "audio": "audio", "video": "video"}.get(kind, "document")


def safe_name(name: str) -> str:
    name = Path(name).name
    return re.sub(r"[^\w.\-؀-ۿ]+", "_", name)[:120] or "file"


def save_upload(fileobj: BinaryIO, directory: Path, max_bytes: int) -> tuple[str, int, str]:
    """ذخیره‌ی جریانی فایل (بدون بار کردن کل آن در حافظه). خروجی: (مسیر موقت، حجم، sha256)."""
    h = hashlib.sha256()
    tmp = tempfile.NamedTemporaryFile(dir=directory, delete=False, suffix=".part")
    size = 0
    try:
        with tmp:
            while chunk := fileobj.read(1024 * 1024):
                size += len(chunk)
                if size > max_bytes:
                    raise db.ValidationError(f"حجم فایل بیش از {max_bytes // (1024 * 1024)} مگابایت است")
                h.update(chunk)
                tmp.write(chunk)
        if size == 0:
            raise db.ValidationError("فایل خالی است")
    except BaseException:
        if os.path.exists(tmp.name):
            os.unlink(tmp.name)
        raise
    return tmp.name, size, h.hexdigest()


def add_document(fileobj: BinaryIO, filename: str, title: str | None = None, category: str | None = None,
                 tags: str | None = None, notes: str | None = None, story_id: int | None = None,
                 doc_date: str | None = None) -> dict[str, Any]:
    kind = kind_of(filename)
    tmp_name, size, digest = save_upload(fileobj, docs_dir(), settings.max_upload_mb * 1024 * 1024)
    try:
        with db.connect() as conn:
            dup = conn.execute("SELECT * FROM documents WHERE sha256 = ?", (digest,)).fetchone()
            if dup:
                os.unlink(tmp_name)
                return {**dict(dup), "duplicate": True}
            if category and category not in CATEGORIES:
                category = None
            if story_id and not conn.execute("SELECT 1 FROM stories WHERE id=?", (story_id,)).fetchone():
                story_id = None
            try:
                doc_date = jalali.normalize(doc_date) if doc_date else None
            except ValueError:
                doc_date = None
            now = db.now_str()
            cur = conn.execute(
                "INSERT INTO documents (title, filename, path, sha256, size, kind, category, tags, notes, doc_date, "
                "story_id, status, created_at, updated_at) VALUES (?, ?, '', ?, ?, ?, ?, ?, ?, ?, ?, 'queued', ?, ?)",
                (title or Path(filename).stem, filename, digest, size, kind, category or default_category(kind),
                 tags, notes, doc_date or db.today_str(), story_id, now, now),
            )
            doc_id = cur.lastrowid
            final = docs_dir() / f"{doc_id}_{safe_name(filename)}"
            os.replace(tmp_name, final)
            conn.execute("UPDATE documents SET path = ? WHERE id = ?", (str(final), doc_id))
            _index_meta(conn, doc_id)
            row = dict(conn.execute("SELECT * FROM documents WHERE id = ?", (doc_id,)).fetchone())
    except BaseException:
        if os.path.exists(tmp_name):
            os.unlink(tmp_name)
        raise
    worker.wake()
    return row


def _index_meta(conn: sqlite3.Connection, doc_id: int) -> None:
    """نمایه‌ی مشخصات سند (عنوان، برچسب، یادداشت) در بند ویژه‌ی صفحه‌ی ۰."""
    d = conn.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
    old = conn.execute("SELECT id FROM doc_chunks WHERE doc_id=? AND page=0", (doc_id,)).fetchall()
    for r in old:
        conn.execute("DELETE FROM doc_fts WHERE chunk_id=?", (r["id"],))
    conn.execute("DELETE FROM doc_chunks WHERE doc_id=? AND page=0", (doc_id,))
    text = "\n".join(x for x in (d["title"], CATEGORIES.get(d["category"] or "", ""), d["tags"], d["notes"],
                                 d["filename"]) if x)
    cur = conn.execute("INSERT INTO doc_chunks (doc_id, page, seq, text, method) VALUES (?, 0, 0, ?, 'meta')",
                       (doc_id, text))
    conn.execute("INSERT INTO doc_fts (norm, chunk_id, doc_id) VALUES (?, ?, ?)",
                 (textnorm.normalize(text), cur.lastrowid, doc_id))


def update_meta(doc_id: int, data: dict[str, Any]) -> dict[str, Any]:
    get_document(doc_id)
    allowed = {"title", "category", "tags", "notes", "story_id", "doc_date"}
    values = {k: (v.strip() if isinstance(v, str) else v) for k, v in data.items() if k in allowed}
    if "title" in values and not values["title"]:
        raise db.ValidationError("عنوان خالی است")
    if values.get("category") and values["category"] not in CATEGORIES:
        raise db.ValidationError("دسته نامعتبر است")
    if "doc_date" in values:
        try:
            values["doc_date"] = jalali.normalize(values["doc_date"]) if values["doc_date"] else None
        except ValueError as e:
            raise db.ValidationError(str(e)) from None
    if "story_id" in values:
        values["story_id"] = int(values["story_id"]) if values["story_id"] not in (None, "") else None
    if values:
        with db.connect() as conn:
            sets = ", ".join(f"{k}=?" for k in values)
            conn.execute(f"UPDATE documents SET {sets}, updated_at=? WHERE id=?",
                         [*values.values(), db.now_str(), doc_id])
            _index_meta(conn, doc_id)
    return get_document(doc_id)


def delete_document(doc_id: int) -> None:
    with db.connect() as conn:
        row = conn.execute("SELECT path FROM documents WHERE id = ?", (doc_id,)).fetchone()
        if not row:
            raise db.NotFound("سند پیدا نشد")
        conn.execute("DELETE FROM doc_fts WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM doc_chunks WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
    if row["path"] and os.path.exists(row["path"]):
        os.unlink(row["path"])


def list_documents(category: str | None = None, story_id: int | None = None) -> list[dict[str, Any]]:
    sql = ("SELECT d.*, s.title AS story_name FROM documents d LEFT JOIN stories s ON s.id = d.story_id")
    where, params = [], []
    if category:
        where.append("d.category = ?")
        params.append(category)
    if story_id:
        where.append("d.story_id = ?")
        params.append(story_id)
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY d.id DESC"
    with db.connect() as conn:
        return [dict(r) for r in conn.execute(sql, params)]


def get_document(doc_id: int) -> dict[str, Any]:
    with db.connect() as conn:
        row = conn.execute("SELECT * FROM documents WHERE id = ?", (doc_id,)).fetchone()
    if not row:
        raise db.NotFound("سند پیدا نشد")
    return dict(row)


def reprocess(doc_id: int) -> None:
    with db.connect() as conn:
        conn.execute("UPDATE documents SET status='queued', error=NULL WHERE id = ?", (doc_id,))
    worker.wake()


def document_text(doc_id: int) -> str:
    with db.connect() as conn:
        rows = conn.execute("SELECT text FROM doc_chunks WHERE doc_id=? AND page>0 ORDER BY page, seq",
                            (doc_id,)).fetchall()
    return "\n\n".join(r["text"] for r in rows)


def render_page_png(doc_id: int, page: int, zoom: float = 1.6) -> bytes:
    import pymupdf

    d = get_document(doc_id)
    if d["kind"] == "image":
        return Path(d["path"]).read_bytes()
    if d["kind"] != "pdf":
        raise db.ValidationError("پیش‌نمایش صفحه فقط برای PDF است")
    with pymupdf.open(d["path"]) as pdf:
        if not 1 <= page <= pdf.page_count:
            raise db.NotFound("صفحه وجود ندارد")
        return pdf[page - 1].get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).tobytes("png")


# ───────────────────────── پردازش (نخ پس‌زمینه) ─────────────────────────


def process_document(doc_id: int) -> None:
    d = get_document(doc_id)
    with db.connect() as conn:
        old = [r["id"] for r in conn.execute("SELECT id FROM doc_chunks WHERE doc_id=? AND page>0", (doc_id,))]
        for cid in old:
            conn.execute("DELETE FROM doc_fts WHERE chunk_id=?", (cid,))
        conn.execute("DELETE FROM doc_chunks WHERE doc_id = ? AND page > 0", (doc_id,))
        conn.execute(
            "UPDATE documents SET status='processing', pages_done=0, ocr_pages=0, fixed_pages=0, weak_pages=0, "
            "chunks=0, error=NULL WHERE id=?", (doc_id,),
        )
        _index_meta(conn, doc_id)
    stats = {"ocr": 0, "fixed": 0, "weak": 0, "chunks": 0}
    pending: list[tuple[int, int, str, str]] = []

    def flush(done: int, total: int) -> None:
        with db.connect() as conn:
            for page, seq, text, method in pending:
                cur = conn.execute(
                    "INSERT INTO doc_chunks (doc_id, page, seq, text, method) VALUES (?, ?, ?, ?, ?)",
                    (doc_id, page, seq, text, method),
                )
                conn.execute("INSERT INTO doc_fts (norm, chunk_id, doc_id) VALUES (?, ?, ?)",
                             (textnorm.normalize(text), cur.lastrowid, doc_id))
            conn.execute(
                "UPDATE documents SET pages=?, pages_done=?, ocr_pages=?, fixed_pages=?, weak_pages=?, chunks=?, "
                "updated_at=? WHERE id=?",
                (total, done, stats["ocr"], stats["fixed"], stats["weak"], stats["chunks"], db.now_str(), doc_id),
            )
        pending.clear()

    total = 0
    for page_no, total, extract in iter_pages(Path(d["path"]), d["kind"]):
        if worker.stopping:
            return
        text, method = extract()
        if method in ("ocr", "fixed", "weak"):
            stats[method] += 1
        for seq, chunk in enumerate(chunk_text(text)):
            pending.append((page_no, seq, chunk, method))
            stats["chunks"] += 1
        if page_no % 5 == 0 or page_no == total:
            flush(page_no, total)
    flush(total, total)
    with db.connect() as conn:
        conn.execute("UPDATE documents SET status='ready', updated_at=? WHERE id=?", (db.now_str(), doc_id))


class Worker:
    def __init__(self) -> None:
        self._event = threading.Event()
        self._thread: threading.Thread | None = None
        self.stopping = False

    def wake(self) -> None:
        self._event.set()

    def start(self) -> None:
        if self._thread and self._thread.is_alive():
            return
        with db.connect() as conn:
            conn.execute("UPDATE documents SET status='queued' WHERE status='processing'")
        self.stopping = False
        self._thread = threading.Thread(target=self._run, name="doc-worker", daemon=True)
        self._thread.start()

    def stop(self) -> None:
        if self._thread and self._thread.is_alive():
            self.stopping = True
            self._event.set()

    def run_pending(self) -> int:
        n = 0
        while not self.stopping:
            with db.connect() as conn:
                row = conn.execute("SELECT id FROM documents WHERE status='queued' ORDER BY id LIMIT 1").fetchone()
            if not row:
                return n
            try:
                with HEAVY:
                    process_document(row["id"])
            except Exception as e:
                log.exception("پردازش سند %s ناموفق بود", row["id"])
                with db.connect() as conn:
                    conn.execute("UPDATE documents SET status='error', error=? WHERE id=?", (str(e)[:500], row["id"]))
            n += 1
        return n

    def _run(self) -> None:
        while not self.stopping:
            try:
                self.run_pending()
            except Exception:
                log.exception("خطای پردازشگر اسناد")
            self._event.wait(30)
            self._event.clear()


worker = Worker()


# ───────────────────────── جستجو ─────────────────────────


def best_sentences(text: str, terms: list[str], n: int = 2) -> list[str]:
    sents = [s.strip() for s in re.split(r"(?<=[.!?؟؛])\s+|\n", text) if len(s.strip()) > 10]
    if not terms:
        return sents[:n]
    scored = []
    for i, s in enumerate(sents):
        norm = textnorm.normalize(s)
        hits = sum(1 for t in terms if t in norm)
        scored.append((hits, -i, s))
    scored.sort(reverse=True)
    return [s for h, _, s in scored[:n] if h > 0] or sents[:1]


def search(query: str, doc_ids: list[int] | None = None, limit: int = 12, category: str | None = None) -> dict[str, Any]:
    """جستجوی کلیدواژه‌ای در متن همه‌ی اسناد؛ بهترین بندها با شماره‌ی صفحه."""
    query = (query or "").strip()
    terms = textnorm.highlight_terms(query)
    if not query:
        return {"query": query, "results": [], "terms": terms}
    results: list[dict[str, Any]] = []
    seen: set[tuple[int, int]] = set()
    with db.connect() as conn:
        # اول همه‌ی واژه‌ها (AND)، اگر نتیجه کم بود هر کدام (OR)
        for mode in ("AND", "OR"):
            fts = textnorm.fts_query(query, mode)
            if not fts:
                break
            sql = ("SELECT f.chunk_id, bm25(doc_fts) AS rank FROM doc_fts f "
                   "JOIN documents d ON d.id = f.doc_id WHERE doc_fts MATCH ?")
            params: list[Any] = [fts]
            if doc_ids:
                sql += f" AND f.doc_id IN ({','.join('?' * len(doc_ids))})"
                params += doc_ids
            if category:
                sql += " AND d.category = ?"
                params.append(category)
            sql += " ORDER BY rank LIMIT ?"
            params.append(limit * 4)
            try:
                ids = [r["chunk_id"] for r in conn.execute(sql, params)]
            except sqlite3.OperationalError as e:
                log.warning("پرسش FTS نامعتبر %r: %s", fts, e)
                ids = []
            if ids:
                rows = {r["id"]: dict(r) for r in conn.execute(
                    f"SELECT c.id, c.doc_id, c.page, c.text, c.method, d.title, d.kind, d.category, d.doc_date "
                    f"FROM doc_chunks c JOIN documents d ON d.id = c.doc_id WHERE c.id IN ({','.join('?' * len(ids))})",
                    ids,
                )}
                for cid in ids:
                    r = rows.get(cid)
                    if not r or (r["doc_id"], r["page"]) in seen:
                        continue
                    seen.add((r["doc_id"], r["page"]))
                    norm = textnorm.normalize(r["text"])
                    results.append({
                        **r, "chunk_id": cid, "all_terms": mode == "AND",
                        "matched_terms": [t for t in terms if t in norm],
                        "highlights": best_sentences(r["text"], terms),
                    })
                    if len(results) >= limit:
                        break
            if len(results) >= limit or len(terms) <= 1:
                break
    return {"query": query, "results": results, "terms": terms}


def stats() -> dict[str, Any]:
    with db.connect() as conn:
        row = conn.execute(
            "SELECT COUNT(*) n, COALESCE(SUM(pages),0) pages, COALESCE(SUM(size),0) size, "
            "COALESCE(SUM(status IN ('queued','processing')),0) pending FROM documents"
        ).fetchone()
    return {**dict(row), "ocr": ocr_available()}
