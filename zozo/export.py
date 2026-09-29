"""خروجی Word راست‌به‌چپ برای متن‌های فارسی (پیاده‌سازی صوت، متن عکس و PDF)."""
from __future__ import annotations

import io


def docx_bytes(title: str, paragraphs: list[str], sections: list[tuple[str, list[str]]] | None = None) -> bytes:
    """یک سند Word فارسی: عنوان، بخش‌های اختیاری (مثلاً خلاصه) و پاراگراف‌های متن."""
    import docx
    from docx.oxml.ns import qn
    from docx.shared import Pt

    d = docx.Document()

    def rtl(text: str, bold: bool = False, size: int = 12) -> None:
        p = d.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.append(pPr.makeelement(qn("w:bidi"), {}))
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        rPr = run._r.get_or_add_rPr()
        rPr.append(rPr.makeelement(qn("w:rtl"), {}))
        run.font.name = "Vazirmatn"
        rPr.rFonts.set(qn("w:cs"), "B Nazanin")

    rtl(title, bold=True, size=16)
    for head, lines in sections or []:
        rtl(head, bold=True, size=13)
        for line in lines:
            if line.strip():
                rtl(line.strip())
    if sections:
        rtl("متن کامل", bold=True, size=13)
    for para in paragraphs:
        if para.strip():
            rtl(para.strip())
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
